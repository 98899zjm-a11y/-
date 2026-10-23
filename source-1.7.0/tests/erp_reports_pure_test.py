"""Pure report and real workbook tests: never imports app or opens SQLite."""
import importlib.util
import sys
import unittest
from pathlib import Path
from io import BytesIO

MODULES=Path(__file__).resolve().parent.parent/'project'/'desktop'/'legacy-source'/'ledger'
sys.path.insert(0,str(MODULES))

class Reports(unittest.TestCase):
    def load(self):
        self.assertTrue((MODULES/'erp_reports.py').is_file(),'The pure ERP report module is not implemented yet')
        spec=importlib.util.spec_from_file_location('erp_reports',MODULES/'erp_reports.py')
        module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module)
        return module
    def actor(self,sensitive=True):
        return {'protocolVersion':2,'staffId':'synthetic-reader','access':'view','capabilities':['erp.read','erp.reports.view']+(['erp.reports.sensitive'] if sensitive else [])}
    def doc(self,id,kind='receivable',date='2026-10-01',status='confirmed',party='CUS:46b24d66-14a8-4ec0-9f89-2b9e10673c1b',amount=10001,payload=None):
        return {'id':id,'number':'ERP-00001','kind':kind,'partyId':party,'partyName':'=HYPERLINK("bad")','date':date,'dueDate':'2026-10-01','status':status,'amountCents':amount,'authorId':'synthetic-maker','authorName':'测试','payload':payload or {},'note':'+1+1','lines':[]}
    def bill(self,id,kind='receivable',amount=10001,settled=1,due='2026-10-01',party='CUS:46b24d66-14a8-4ec0-9f89-2b9e10673c1b'):
        return {'id':id,'number':'AR-00001','kind':kind,'partyId':party,'partyName':'=HYPERLINK("bad")','date':'2026-10-01','dueDate':due,'amountCents':amount,'settledCents':settled,'balanceCents':amount-settled,'sourceId':id}
    def test_strict_filters_and_calendar_ranges(self):
        m=self.load()
        valid=m.parse_filters({'asOf':'2026-10-14','from':'2026-10-01','to':'2026-10-10','partyId':'SUP:2','kind':'receipt','status':'confirmed','creatorId':'synthetic-maker'})
        self.assertEqual(valid['asOf'],'2026-10-14')
        for params in ({'asOf':'2026-02-30'},{'asOf':'2026-10-14','from':'2026-10-15'},{'asOf':'2026-10-14','from':'2026-10-05','to':'2026-10-01'},{'kind':'any-sql'},{'status':'deleted'},{'actor':'admin'},{'partyId':'../secret'},{'creatorId':'a\nadmin'}):
            with self.subTest(params=params),self.assertRaises(ValueError):m.parse_filters(params)
    def test_aging_and_totals_use_integer_cents_and_due_day_boundaries(self):
        m=self.load();docs=[self.doc('a'),self.doc('b')];bills=[self.bill('a',due='2026-10-14'),self.bill('b',amount=102,settled=2,due='2026-09-14')]
        r=m.build_report(self.actor(),{'asOf':'2026-10-14'},bills,docs)
        self.assertEqual(r['totals']['receivableCents'],10100);self.assertEqual(r['totals']['receivableCount'],2)
        self.assertEqual(r['aging'][0]['receivableCents'],10000);self.assertEqual(r['aging'][1]['receivableCents'],100)
        self.assertEqual(sum(x['receivableCents'] for x in r['aging']),r['totals']['receivableCents'])
    def test_no_sensitive_costs_in_nested_reports_or_export(self):
        m=self.load();docs=[self.doc('ar'),self.doc('ap','receipt',party='SUP:2',amount=987654321),self.doc('pay','settlement',party='SUP:2',payload={'direction':'pay','account':'secret-bank','amountCents':987654321}),self.doc('order','purchase',party='SUP:2',amount=987654321)]
        r=m.build_report(self.actor(False),{'asOf':'2026-10-14'},[self.bill('ar'),self.bill('ap','payable',987654321,party='SUP:2')],docs)
        self.assertEqual([x['id'] for x in r['rows']],['ar']);self.assertEqual(r['cash'],[]);self.assertEqual(r['purchase'],[])
        self.assertNotIn('payableCents',r['totals']);self.assertNotIn('paidCents',r['totals']);self.assertTrue(all('payableCents' not in x for x in r['aging']))
        self.assertNotIn('987654321',str(r));self.assertNotIn('secret-bank',str(r))
    def test_cash_reversals_and_current_status_do_not_rewrite_historical_asof(self):
        m=self.load();original=self.doc('cash','settlement',date='2026-10-02',status='reversed',amount=30,payload={'direction':'receive','account':'虚构现金'})
        reversal=self.doc('undo','reversal',date='2026-10-10',amount=30,payload={'targetId':'cash'})
        early=m.build_report(self.actor(),{'asOf':'2026-10-05'},[],[original,reversal]);late=m.build_report(self.actor(),{'asOf':'2026-10-14'},[],[original,reversal])
        self.assertEqual(early['totals']['receivedCents'],30);self.assertEqual(late['totals']['receivedCents'],0)
        self.assertEqual([x['amountCents'] for x in late['cash']],[30,-30]);self.assertEqual(early['cash'][0]['status'],'confirmed')
    def test_source_filters_do_not_recalculate_asof_balance_from_cash_window(self):
        m=self.load();docs=[self.doc('a'),self.doc('b','opening_receivable',date='2026-10-02')]
        r=m.build_report(self.actor(),{'asOf':'2026-10-14','kind':'receivable','creatorId':'synthetic-maker'},[self.bill('a',settled=2500),self.bill('b',settled=400)],docs)
        self.assertEqual([x['id'] for x in r['rows']],['a']);self.assertEqual(r['rows'][0]['balanceCents'],7501)
        self.assertEqual(r['totals']['receivableCents'],7501)
    def test_real_xlsx_roundtrip_keeps_untrusted_text_nonformula_and_money_exact(self):
        m=self.load();r=m.build_report(self.actor(),{'asOf':'2026-10-14'},[self.bill('a')],[self.doc('a')])
        output=m.workbook(r);self.assertIsInstance(output,BytesIO);self.assertEqual(output.getvalue()[:2],b'PK')
        from openpyxl import load_workbook
        book=load_workbook(output,data_only=False,read_only=True)
        try:
            self.assertEqual(book.sheetnames,['往来余额','汇总','账龄','收付款','采购验收','筛选条件'])
            self.assertTrue(all(cell.data_type!='f' for sheet in book for row in sheet.iter_rows() for cell in row))
            cells=[cell.value for row in book.active.iter_rows() for cell in row]
            self.assertIn('=HYPERLINK("bad")',cells);self.assertIn('100.01',cells);self.assertIn('0.01',cells);self.assertIn('100.00',cells)
        finally:book.close()
    def test_large_valid_bills_cannot_round_the_browser_portfolio_total(self):
        m=self.load();docs=[];bills=[]
        for index in range(1001):
            amount=9000000000000 if index<1000 else 8000000000001
            key='precision-'+str(index);docs.append(self.doc(key,amount=amount));bills.append(self.bill(key,amount=amount,settled=0))
        with self.assertRaises(ValueError):m.build_report(self.actor(),{'asOf':'2026-10-14'},bills,docs)
    def test_unrecognized_reversal_target_is_hidden_without_recursive_traversal(self):
        m=self.load();bad=self.doc('cycle','reversal',payload={'targetId':'cycle'})
        r=m.build_report(self.actor(False),{'asOf':'2026-10-14'},[],[bad])
        self.assertEqual(r['cash'],[]);self.assertEqual(r['purchase'],[])

if __name__=='__main__':unittest.main()
