from pathlib import Path
ROOT = Path(__file__).resolve().parent.parent
from python_guard import install
install(ROOT, require_restricted=True)
from openpyxl import load_workbook
book=load_workbook(ROOT/'data/documents/export-fixture.xlsx',read_only=True,data_only=False)
assert book.sheetnames==['报价历史']
rows=list(book.active.iter_rows(values_only=True))
assert rows and rows[0][0]=='报价单号'
assert all(cell.data_type!='f' for row in book.active for cell in row)
book.close()
print('Real OOXML workbook read successfully with openpyxl; no formula cells.')
