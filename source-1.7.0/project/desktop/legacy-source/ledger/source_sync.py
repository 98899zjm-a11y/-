"""Original rows retained verbatim; business records linked to stable file/sheet/row IDs."""
import pathlib,json,hashlib,datetime,collections,math,re
import openpyxl
from app import importer,db
import repairs as r

SCHEMA='''
CREATE TABLE IF NOT EXISTS source_documents(id INTEGER PRIMARY KEY,name TEXT NOT NULL,sha256 TEXT UNIQUE NOT NULL,path TEXT NOT NULL,created_at TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS source_rows(id INTEGER PRIMARY KEY,document_id INTEGER NOT NULL,sheet TEXT NOT NULL,row_no INTEGER NOT NULL,kind TEXT NOT NULL,raw_json TEXT NOT NULL,formula_json TEXT NOT NULL,normalized_json TEXT NOT NULL,classification TEXT NOT NULL,warning TEXT NOT NULL,UNIQUE(document_id,sheet,row_no));
CREATE TABLE IF NOT EXISTS source_links(id INTEGER PRIMARY KEY,source_row_id INTEGER NOT NULL,part TEXT NOT NULL,ledger_table TEXT NOT NULL,ledger_id INTEGER NOT NULL,expected_json TEXT NOT NULL,UNIQUE(source_row_id,part));
CREATE INDEX IF NOT EXISTS idx_source_ledger ON source_links(ledger_table,ledger_id);
'''
TITLES={'sales':'销售','stock':'库存','material':'存料费','other':'说明/未识别'}
SALE_HEADERS={**importer.SALES_MAP,**importer.SALES_CHECK,'labor':('工费','加工费'),'spread':('点差',),
              'gross_rate_c':('毛利率',),'comm_c':('提成','销售提成')}
STOCK_HEADERS={**importer.STOCK_MAP,**importer.STOCK_CHECK,'note':('备注2','备注二')}
TEXT_SALES=set(r.SALE_COLUMNS)-r.SALE_NUMBERS
NUMERIC=set(r.SALE_NUMBERS)|set(r.SALE_DERIVED)|{'sell_price','profit_c','seq','grams','amount'}
DATES={'date','remit_date','in_date','sell_date'}

def j(value):return json.dumps(value,ensure_ascii=False,default=lambda x:x.isoformat() if isinstance(x,(datetime.date,datetime.datetime)) else str(x))
def norm(x):return importer._norm(x)
def numeric(x):
    if x in (None,''):return None
    if isinstance(x,bool):return int(x)
    try:value=float(str(x).replace(',','').replace('￥','').replace('¥','').replace('元','').replace('克','').strip())
    except ValueError:raise ValueError('不能识别数字：'+str(x))
    if not math.isfinite(value):raise ValueError('不能导入无限值。')
    return value

def map_headers(header,mapping):
    h=[norm(x) for x in header];return {field:next(i for i,v in enumerate(h) if v in [norm(a) for a in aliases]) for field,aliases in mapping.items() if any(v in [norm(a) for a in aliases] for v in h)}

def material_summary(raw):
    """Summary labels describe aggregate rows, regardless of their row number.

    Check the date/measure cells only; supplier names or comments can contain
    these words without making an ordinary material transaction a summary.
    """
    labels=[norm(value) for value in list(raw)[:2] if isinstance(value,str)]
    return any(re.fullmatch(r'(?:料|费|板料|金额|克数|重量|本月|本期|当月|月度|累计|总)?(?:合计|总计|小计|汇总)(?:[（(].*[）)])?',label) for label in labels)

def material_unit(value,number_format='',measure=''):
    """The operator authorized unmarked material figures to mean grams.

    An explicit amount row or a value/format carrying 元 denotes money; formats
    are otherwise presentation and percentages do not imply a material unit.
    """
    if '金额' in str(measure) or '元' in str(value) or '元' in str(number_format):return 'amount'
    return 'grams'

def confirmed_blank_sell_date(sheet,row_no,value):
    """Two original text cells were explicitly confirmed as undated by the user."""
    return sheet=='黄白铂板料9月' and row_no in (21,43) and str(value).strip()=='万德福'

def parse_book(path):
    path=pathlib.Path(path)
    values=openpyxl.load_workbook(path,data_only=True);formulas=openpyxl.load_workbook(path,data_only=False)
    rows=[];info=[]
    for ws in values:
        allrows=list(ws.values);header_index=None;kind='other';mapping={};suppliers=[]
        for idx,row in enumerate(allrows[:20]):
            h=[norm(v) for v in row]
            if '日期' in h and len(h)>1 and h[1]=='商品' and '销售克价' not in h and '成本' not in h:
                header_index=idx;kind='material';suppliers=[(i,str(v).strip()) for i,v in enumerate(row) if i>1 and v and str(v) not in ['汇总','备注'] and not str(v).startswith('列')];break
            if '商品' in h and ('销售克价' in h or '实际到账' in h):header_index=idx;kind='sales';mapping=map_headers(row,SALE_HEADERS);mapping.setdefault('date',0);break
            if '商品' in h and '成本' in h:header_index=idx;kind='stock';mapping=map_headers(row,STOCK_HEADERS);break
        if kind=='stock' and header_index is not None:
            notes=[i for i,v in enumerate(allrows[header_index]) if norm(v)=='备注']
            if len(notes)>1:mapping['note']=notes[1]
        counts=collections.Counter();lastdate=''
        for idx,raw in enumerate(allrows,1):
            raw=list(raw)
            while raw and raw[-1] is None:raw.pop()
            if not any(v not in (None,'') for v in raw):classification='空行';normalized={};warning=''
            elif header_index is None or idx<=header_index+1:classification='表头/说明';normalized={};warning=''
            else:
                classification='业务明细';normalized={};warning=''
                if kind in ('sales','stock'):
                    for field,col in mapping.items():
                        value=raw[col] if col<len(raw) else None
                        if value in (None,''):continue
                        try:
                            if field in NUMERIC:normalized[field]=numeric(value)
                            elif field in DATES:
                                if field=='sell_date' and confirmed_blank_sell_date(ws.title,idx,value):
                                    normalized[field]='';continue
                                parsed=importer.ymd(value)
                                if not parsed:raise ValueError('日期不能识别：'+str(value))
                                normalized[field]=parsed
                            else:normalized[field]=str(value).strip()
                        except ValueError as e:warning+=str(e)+'；'
                    product=normalized.get('product','')
                    if product in ('商品','合计','总计','小计') or not product:
                        classification='说明/汇总'
                        if any(isinstance(v,(int,float)) and v for v in raw) and product not in ('商品','合计','总计','小计'):
                            warning='这一行有数字但缺少商品，保留原行待核对。';classification='待核对'
                    elif kind=='sales' and not normalized.get('date'):
                        warning+='缺少销售日期，不自动补今天。';classification='待核对'
                    if warning and classification=='业务明细':classification='待核对'
                elif kind=='material':
                    date=importer.ymd(raw[0]) if raw else ''
                    if date:lastdate=date
                    measure=str(raw[1] or '') if len(raw)>1 else ''
                    numeric_cells=[]
                    for col,supplier in suppliers:
                        value=raw[col] if col<len(raw) else None
                        if value in (None,''):continue
                        try:
                            n=numeric(value);fmt=formulas[ws.title].cell(idx,col+1).number_format
                            numeric_cells.append({'supplier':supplier,'value':n,'column':col+1,'unit':material_unit(value,fmt,measure),'number_format':fmt})
                        except ValueError as e:warning+=str(e)+'；'
                    if material_summary(raw):classification='说明/汇总';normalized={}
                    elif not numeric_cells:classification='说明/空格';normalized={}
                    elif not lastdate:classification='待核对';warning+='缺少日期。'
                    else:
                        unit='amount' if '金额' in measure else 'grams'
                        normalized={'date':lastdate,'unit':unit,'cells':numeric_cells}
                else:classification='未识别';warning='没有识别业务表头，原行仍保留。'
            fm={openpyxl.utils.get_column_letter(c.column):c.value for c in formulas[ws.title][idx] if c.data_type=='f'}
            if kind=='stock' and len(raw)>9 and confirmed_blank_sell_date(ws.title,idx,raw[9]):
                warning+='操作员已确认原销售时间文字“万德福”的日期留空；商家按原M列保留。'
            for c in values[ws.title][idx]:
                if c.data_type=='e':warning+='原表 '+c.coordinate+' 为 '+str(c.value)+'；'
            counts[classification]+=1
            rows.append(dict(sheet=ws.title,row_no=idx,kind=kind,raw=raw,formulas=fm,normalized=normalized,classification=classification,warning=warning))
        info.append({'sheet':ws.title,'kind':kind,'rows':ws.max_row,'counts':dict(counts),'mapping':{k:openpyxl.utils.get_column_letter(v+1) for k,v in mapping.items()}})
    values.close();formulas.close();return rows,info

def register(path,name,c):
    raw=pathlib.Path(path).read_bytes();sha=hashlib.sha256(raw).hexdigest()
    old=c.execute('SELECT id FROM source_documents WHERE sha256=?',(sha,)).fetchone()
    if old:return old[0]
    rows,info=parse_book(path)
    doc=c.execute('INSERT INTO source_documents(name,sha256,path,created_at) VALUES(?,?,?,?)',(name,sha,str(path),db.now_str())).lastrowid
    for x in rows:
        c.execute('INSERT INTO source_rows(document_id,sheet,row_no,kind,raw_json,formula_json,normalized_json,classification,warning) VALUES(?,?,?,?,?,?,?,?,?)',
                  (doc,x['sheet'],x['row_no'],x['kind'],j(x['raw']),j(x['formulas']),j(x['normalized']),x['classification'],x['warning']))
    return doc

def ledger_values(kind,normalized):
    if kind=='sales':
        rec={k:normalized.get(k,0.0 if k in r.SALE_NUMBERS else '') for k in r.SALE_COLUMNS}
        rec['status']=rec['status'] or '已完成'
        if 'labor' not in normalized and 'amount_c' in normalized:rec['labor']=normalized['amount_c']-rec['sell_gram_price']*rec['weight']
        if 'spread' not in normalized and 'cost_gram_c' in normalized:rec['spread']=rec['sell_gram_price']-normalized['cost_gram_c']
        if 'spread' not in normalized and 'cost_gram_c' not in normalized:rec['spread']=8.2 if r.category_spread(rec)=='spread_bullion' else 40.0
        calc_input=dict(rec,_rule_id=0)
        for target,src in r.SALE_DERIVED.items():
            if target in normalized:calc_input[src]=normalized[target]
        # Imported source values retain full numeric precision; no money rounding.
        values=r.calculate_sale(calc_input)
        rec.update({key:normalized[key] if key in normalized else values[src] for key,src in r.SALE_DERIVED.items()})
        if 'comm_c' not in normalized:
            if 'profit_c' in normalized and 'comm_profit_c' in normalized:rec['comm_c']=normalized['profit_c']-normalized['comm_profit_c']
            elif 'comm_cost_c' in normalized:rec['comm_c']=normalized['comm_cost_c']-rec['cost']
        rec['rule_version']=0;return rec
    rec={k:normalized.get(k,0.0 if k in ['weight','cost','sell_price'] else '') for k in r.stock.FIELDS}
    rec['profit_c']=normalized.get('profit_c',0.0);rec['rule_version']=0
    if 'seq' in normalized:rec['seq']=int(normalized['seq'])
    return rec

def import_document(doc,c,decisions=None,selected=None):
    decisions=decisions or {};counts=collections.Counter();material_atoms=[];legacy_formats=None
    rows=c.execute('SELECT * FROM source_rows WHERE document_id=? ORDER BY id',(doc,)).fetchall()
    for source in rows:
        source=dict(source);kind=source['kind'];normalized=json.loads(source['normalized_json'])
        if selected is not None and source['sheet'] not in selected:continue
        raw=json.loads(source['raw_json'])
        if kind=='stock' and len(raw)>9 and confirmed_blank_sell_date(source['sheet'],source['row_no'],raw[9]):
            normalized['sell_date']=''
            remaining=source['warning'].replace('日期不能识别：万德福；','').strip('； ')
            if source['classification']=='待核对' and not remaining:source['classification']='业务明细'
            confirmation='操作员已确认原销售时间文字“万德福”的日期留空；商家按原M列保留。'
            source['warning']=remaining if confirmation in remaining else (remaining+'；' if remaining else '')+confirmation
            c.execute('UPDATE source_rows SET normalized_json=?,classification=?,warning=? WHERE id=?',(j(normalized),source['classification'],source['warning'],source['id']))
        # Apply the same protection to documents registered by earlier versions.
        if kind=='material' and material_summary(json.loads(source['raw_json'])):
            if source['classification']!='说明/汇总':
                c.execute('UPDATE source_rows SET classification=?,normalized_json=?,warning=? WHERE id=?',('说明/汇总',j({}),'原表汇总行保留，不作为存料明细导入。',source['id']))
            counts['非业务/汇总']+=1;continue
        if kind=='material' and normalized.get('cells'):
            raw=json.loads(source['raw_json']);measure=raw[1] if len(raw)>1 else ''
            for cell in normalized['cells']:
                choice=decisions.get(str(source['id'])+'_'+str(cell['column']),decisions.get(str(source['id']),''))
                if choice in ['grams','amount','ignore']:cell['unit']=choice
                elif not cell.get('unit'):
                    if 'number_format' not in cell and legacy_formats is None:
                        legacy_formats={};record=c.execute('SELECT path FROM source_documents WHERE id=?',(doc,)).fetchone()
                        path=pathlib.Path(record['path']) if record else None
                        if path and not path.is_absolute():path=r.config.data_dir()/path
                        if path and path.is_file():
                            legacy_book=openpyxl.load_workbook(path,data_only=True,read_only=True)
                            try:
                                sheets={x['sheet'] for x in rows if x['kind']=='material'}
                                for sheet in sheets:
                                    if sheet not in legacy_book.sheetnames:continue
                                    for original_row in legacy_book[sheet].iter_rows():
                                        for original_cell in original_row:
                                            if original_cell.value is not None:legacy_formats[(sheet,original_cell.row,original_cell.column)]=original_cell.number_format
                            finally:legacy_book.close()
                    fmt=cell.get('number_format',(legacy_formats or {}).get((source['sheet'],source['row_no'],cell['column']),''))
                    value=raw[cell['column']-1] if cell['column']<=len(raw) else None
                    cell['unit']=material_unit(value,fmt,measure);cell['number_format']=fmt
            normalized['unit']='amount' if '金额' in str(measure) else 'grams'
            if source['classification']=='待选择单位':
                source['classification']='业务明细';source['warning']='按操作员规则：未标单位默认克数，明确元按金额；原行保留。'
            c.execute('UPDATE source_rows SET normalized_json=?,classification=?,warning=? WHERE id=?',(j(normalized),source['classification'],source['warning'],source['id']))
        linked_parts={x[0] for x in c.execute('SELECT part FROM source_links WHERE source_row_id=?',(source['id'],))}
        if linked_parts and kind!='material':counts['已对应，跳过']+=1;continue
        classification=source['classification']
        if classification=='待选择单位':
            for cell in normalized['cells']:
                choice=decisions.get(str(source['id'])+'_'+str(cell['column']),decisions.get(str(source['id']),''))
                if choice in ['grams','amount','ignore']:cell['unit']=choice
            complete=all(cell.get('unit') for cell in normalized['cells'])
            c.execute('UPDATE source_rows SET normalized_json=?,classification=?,warning=? WHERE id=?',(j(normalized),'业务明细' if complete else classification,'操作员指定单元格单位；未指定的格不会导入。',source['id']))
            classification='业务明细'
        if classification!='业务明细':counts['非业务/待核对']+=1;continue
        if kind in ['sales','stock']:
            table='sales' if kind=='sales' else 'stock_items';values=ledger_values(kind,normalized)
            expected={k:v for k,v in normalized.items() if k in values}
            # Upgrade safety: reuse an unambiguously identical existing record instead
            # of appending a second copy of the old imported ledger.
            clause='product=? AND order_no=?';params=[values['product'],values['order_no']]
            if kind=='sales' and normalized.get('date'):clause+=' AND date=?';params.append(normalized['date'])
            candidates=[dict(x) for x in c.execute('SELECT * FROM '+table+' WHERE '+clause,params)]
            available=[]
            for candidate in candidates:
                claimed=c.execute('SELECT l.id FROM source_links l JOIN source_rows s ON s.id=l.source_row_id WHERE l.ledger_table=? AND l.ledger_id=? AND s.document_id=? LIMIT 1',(table,candidate['id'],doc)).fetchone()
                if not claimed:available.append(candidate)
            def same(candidate):
                for field,value in expected.items():
                    actual=candidate[field]
                    if isinstance(value,(int,float)) and isinstance(actual,(int,float)):
                        if abs(value-actual)>1e-7:return False
                    elif value!=actual:return False
                return True
            identical=[x for x in available if same(x)]
            if len(identical)==1:
                rid=identical[0]['id']
                c.execute('INSERT INTO source_links(source_row_id,part,ledger_table,ledger_id,expected_json) VALUES(?,?,?,?,?)',(source['id'],'record',table,rid,j(expected)))
                counts['已匹配旧账']+=1;continue
            if available:
                c.execute('UPDATE source_rows SET classification=?,warning=? WHERE id=?',('现有账目待匹配','系统已有相同单号/商品，但不能唯一逐字段对应。请人工匹配，未重复入账。',source['id']))
                counts['旧账待匹配']+=1;continue
            for table_name,field in [('staff','staff'),('channels','channel'),('categories','category'),('suppliers','supplier')]:
                if values.get(field):c.execute('INSERT OR IGNORE INTO '+table_name+'(name,sort,active) VALUES(?,999,1)',(values[field],))
            rid=r.write_row(c,table,values)
            c.execute('INSERT INTO source_links(source_row_id,part,ledger_table,ledger_id,expected_json) VALUES(?,?,?,?,?)',(source['id'],'record',table,rid,j(expected)))
            counts[kind]+=1
        elif kind=='material':
            for cell in normalized['cells']:
                if str(cell['column']) in linked_parts:continue
                unit=cell.get('unit') or normalized.get('unit')
                if unit in ['grams','amount']:material_atoms.append((source,normalized['date'],unit,cell))
                elif unit=='ignore':counts['忽略单元格']+=1
                else:counts['待选择单位单元格']+=1
    groups=collections.defaultdict(list)
    for atom in material_atoms:groups[(atom[1],atom[3]['supplier'])].append(atom)
    for (date,supplier),atoms in groups.items():
        # Never silently merge imports into an existing manually recorded day/supplier.
        matches=c.execute('SELECT * FROM material_stock WHERE date=? AND supplier=?',(date,supplier)).fetchall()
        existing=dict(matches[0]) if len(matches)==1 else None
        owners={x[0] for x in c.execute('SELECT s.document_id FROM source_links l JOIN source_rows s ON s.id=l.source_row_id WHERE l.ledger_table=? AND l.ledger_id=?',('material_stock',matches[0]['id']))} if matches else set()
        c.execute('INSERT OR IGNORE INTO suppliers(name,sort,active) VALUES(?,999,1)',(supplier,))
        values={'date':date,'supplier':supplier,'grams':sum(a[3]['value'] for a in atoms if a[2]=='grams'),
                'amount':sum(a[3]['value'] for a in atoms if a[2]=='amount'),'note':'原表导入'}
        if matches and (len(matches)!=1 or owners!={doc}):
            same_values=len(matches)==1 and abs(existing['grams']-values['grams'])<=1e-7 and abs(existing['amount']-values['amount'])<=1e-7
            if same_values:
                rid=existing['id'];counts['已匹配旧存料']+=1
            else:
                for source,_date,unit,cell in atoms:
                    stored=c.execute('SELECT normalized_json FROM source_rows WHERE id=?',(source['id'],)).fetchone()
                    norm_values=json.loads(stored[0])
                    for stored_cell in norm_values['cells']:
                        if stored_cell['column']==cell['column']:
                            stored_cell['existing_ids']=[x['id'] for x in matches]
                            stored_cell['warning']='同一天同一家已有存料记录，金额或克数不能直接对应，未重复入账。'
                    c.execute('UPDATE source_rows SET normalized_json=?,warning=? WHERE id=?',(j(norm_values),supplier+' 旧存料待人工匹配；其他明确格可继续导入。',source['id']))
                counts['旧存料待匹配']+=1;continue
        else:
            if existing:
                values['grams']+=existing['grams'];values['amount']+=existing['amount']
            rid=r.write_row(c,'material_stock',values,existing)
        for source,_date,unit,cell in atoms:
            expected={'date':date,'supplier':supplier,unit:cell['value']}
            c.execute('INSERT INTO source_links(source_row_id,part,ledger_table,ledger_id,expected_json) VALUES(?,?,?,?,?)',(source['id'],str(cell['column']), 'material_stock',rid,j(expected)))
        counts['material']+=1
    return dict(counts)

def reconcile(doc,c):
    rows=c.execute('SELECT * FROM source_rows WHERE document_id=? ORDER BY id',(doc,)).fetchall()
    links=[dict(x) for x in c.execute('SELECT l.*,s.sheet,s.row_no FROM source_links l JOIN source_rows s ON s.id=l.source_row_id WHERE s.document_id=?',(doc,))]
    groups=collections.defaultdict(list)
    for link in links:groups[(link['ledger_table'],link['ledger_id'])].append(link)
    items=[]
    for (table,rid),group in groups.items():
        ledger=c.execute('SELECT * FROM '+table+' WHERE id=?',(rid,)).fetchone()
        expected={};addresses=[]
        for link in group:
            addresses.append(link['sheet']+'!第'+str(link['row_no'])+'行'+('／第'+link['part']+'列' if link['part']!='record' else ''))
            for field,value in json.loads(link['expected_json']).items():
                if table=='material_stock' and field in ['grams','amount']:expected[field]=expected.get(field,0)+value
                else:expected[field]=value
        diffs=[]
        if ledger:
            for field,value in expected.items():
                actual=ledger[field]
                equal=abs(value-actual)<=1e-7 if isinstance(value,(int,float)) and isinstance(actual,(int,float)) else value==actual
                if not equal:diffs.append({'field':field,'source':value,'system':actual})
        else:diffs.append({'field':'记录','source':'存在','system':'已删除'})
        items.append(dict(table=table,id=rid,addresses=addresses,differences=diffs,matched=not diffs))
    linked={x['source_row_id'] for x in links}
    parts=collections.defaultdict(set)
    for link in links:parts[link['source_row_id']].add(link['part'])
    pending=[]
    for row in rows:
        remaining=False
        if row['kind']=='material':
            value=json.loads(row['normalized_json'])
            remaining=any(str(cell['column']) not in parts[row['id']] and cell.get('unit')!='ignore' for cell in value.get('cells',[]))
        if remaining or row['classification']=='待选择单位' or row['id'] not in linked and row['classification'] in ['业务明细','待核对','未识别','现有账目待匹配']:
            pending.append(dict(id=row['id'],sheet=row['sheet'],row_no=row['row_no'],classification=row['classification'],warning=row['warning']))
    system_extra={table:c.execute('SELECT COUNT(*) FROM '+table+' WHERE id NOT IN (SELECT ledger_id FROM source_links WHERE ledger_table=?)',(table,)).fetchone()[0] for table in ['sales','stock_items','material_stock']}
    return dict(items=items,matched=sum(x['matched'] for x in items),differences=sum(not x['matched'] for x in items),pending=pending,source_rows=len(rows),linked_rows=len(linked),system_extra=system_extra)
