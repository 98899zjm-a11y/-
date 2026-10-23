"""Three-way synchronization of a dedicated Excel working copy.

Archived uploads are immutable. File I/O is performed outside SQLite transactions.
The shared workbook editor remains the only path that changes internal records.
"""
import collections
import contextlib
import datetime
import hashlib
import io
import json
import os
import pathlib
import threading
import uuid

import openpyxl
from flask import Blueprint, abort, flash, redirect, render_template, request, url_for
from app import auth, config, db
import repairs as r
import enhancements as enhanced
import workbooks as wb


bp = Blueprint('external_sync', __name__)
SCHEMA = '''
CREATE TABLE IF NOT EXISTS external_workbook_sync(
 document_id INTEGER PRIMARY KEY,
 enabled INTEGER NOT NULL DEFAULT 0,
 status TEXT NOT NULL DEFAULT 'paused',
 message TEXT NOT NULL DEFAULT '',
 baseline_json TEXT NOT NULL DEFAULT '{}',
 file_sha TEXT NOT NULL DEFAULT '',
 revision INTEGER NOT NULL DEFAULT 0,
 generation INTEGER NOT NULL DEFAULT 0,
 conflicts_json TEXT NOT NULL DEFAULT '[]',
 updated_at TEXT NOT NULL DEFAULT ''
);
'''
STATUS_LABELS = {
    'active': '自动同步已开启', 'waiting_file': '正在等待 Excel 保存稳定',
    'waiting_write': '系统已更新，等待回写 Excel', 'paused': '已暂停',
    'conflict': '双方修改冲突，已暂停', 'structure': '原表结构变化，已暂停',
    'error': '需要处理，已暂停',
}
_locks = {}
_locks_guard = threading.Lock()
_observations = {}
_started = set()


class FileChanged(ValueError):
    pass


class StructureChanged(ValueError):
    pass


def document_lock(doc):
    with _locks_guard:
        return _locks.setdefault(int(doc), threading.RLock())


def dump(value):
    return json.dumps(value, ensure_ascii=False, allow_nan=False)


def cell_key(sheet, row, col):
    return json.dumps([sheet, int(row), int(col)], ensure_ascii=False, separators=(',', ':'))


def normalize(value):
    if isinstance(value, datetime.datetime):
        return value.date().isoformat() if value.time() == datetime.time() else value.isoformat()
    if isinstance(value, datetime.date):
        return value.isoformat()
    if isinstance(value, datetime.time):
        return value.isoformat()
    if value == '':
        return None
    if isinstance(value, str) and value.endswith('T00:00:00'):
        try:
            return datetime.date.fromisoformat(value[:10]).isoformat()
        except ValueError:
            pass
    return value


def entry(value=None, formula=''):
    return {'value': normalize(value), 'formula': formula or ''}


def same(left, right):
    left, right = left or entry(), right or entry()
    if left.get('formula') or right.get('formula'):
        return left.get('formula', '') == right.get('formula', '')
    if isinstance(left.get('value'), bool) != isinstance(right.get('value'), bool):
        return False
    return left.get('value') == right.get('value')


def input_value(value):
    if value.get('formula'):
        return value['formula']
    return '' if value.get('value') is None else value['value']


def snapshot(rows, cells, formulas):
    limits = collections.OrderedDict()
    for row in rows:
        meta = limits.setdefault(row['sheet'], {'rows': [], 'columns': 1})
        meta['rows'].append(row['row_no'])
        meta['columns'] = max(meta['columns'], len(row['cells']),
                              max([int(col) for col in row['formulas']] + [0]))
    values = {}
    for sheet, meta in limits.items():
        for row in meta['rows']:
            for col in range(1, meta['columns'] + 1):
                key = (sheet, row, col)
                values[cell_key(*key)] = entry(cells.get(key), formulas.get(key, ''))
    return {'entries': values, 'limits': limits}


def safe_path(record):
    """Only a doc-id filename directly inside data/联动表格 is permitted."""
    data_root = pathlib.Path(config.data_dir()).resolve()
    folder = (data_root / '联动表格').resolve()
    if folder.parent != data_root:
        raise ValueError('联动目录必须位于当前 data 文件夹内。')
    suffix = '.xlsm' if pathlib.Path(record['name']).suffix.lower() == '.xlsm' else '.xlsx'
    filename = '原表_' + str(int(record['id'])) + suffix
    path = (folder / filename).resolve()
    if path.parent != folder or path.name != filename:
        raise ValueError('联动文件路径不在允许的目录内。')
    return path


def read_bytes(path):
    data = path.read_bytes()
    return data, hashlib.sha256(data).hexdigest()


def atomic_write(path, data, expected_sha):
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = (path.parent / ('.' + path.name + '.' + uuid.uuid4().hex + '.tmp')).resolve()
    if temporary.parent != path.parent.resolve():
        raise ValueError('临时文件超出联动目录。')
    try:
        with temporary.open('xb') as output:
            output.write(data)
            output.flush()
            os.fsync(output.fileno())
        actual_sha = read_bytes(path)[1] if path.exists() else None
        if actual_sha != expected_sha:
            raise FileChanged('Excel 文件刚刚再次保存，等待重新读取，未覆盖文件。')
        os.replace(temporary, path)
    finally:
        with contextlib.suppress(OSError):
            temporary.unlink()


def backup_copy(path, data):
    name = path.stem + '_处理前_' + datetime.datetime.now().strftime('%Y%m%d_%H%M%S') + '_' + uuid.uuid4().hex[:6] + path.suffix
    target = (path.parent / name).resolve()
    if target.parent != path.parent.resolve():
        raise ValueError('备份路径超出联动目录。')
    with target.open('xb') as output:
        output.write(data)
    return target.name


def read_external(data, limits):
    formulas = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    cached = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        layout = {'sheets': formulas.sheetnames,
                  'dimensions': {sheet.title: [sheet.max_row, sheet.max_column] for sheet in formulas}}
        values, outside = {}, {}
        for sheet in formulas:
            meta = limits.get(sheet.title, {'rows': [], 'columns': 0})
            allowed_rows = set(meta['rows'])
            cached_sheet = cached[sheet.title]
            for formula_row, value_row in zip(sheet.iter_rows(), cached_sheet.iter_rows()):
                for formula_cell, value_cell in zip(formula_row, value_row):
                    if formula_cell.value is None and value_cell.value is None:
                        continue
                    row, col = formula_cell.row, formula_cell.column
                    formula = str(formula_cell.value) if formula_cell.data_type == 'f' else ''
                    current = entry(value_cell.value if formula else formula_cell.value, formula)
                    key = cell_key(sheet.title, row, col)
                    if row in allowed_rows and col <= meta['columns']:
                        values[key] = current
                    else:
                        outside[key] = current
        # Blank cells remain explicit so clearing a value is an actual change.
        for sheet, meta in limits.items():
            for row in meta['rows']:
                for col in range(1, meta['columns'] + 1):
                    values.setdefault(cell_key(sheet, row, col), entry())
        return {'entries': values, 'limits': limits, 'layout': layout, 'outside': outside}
    finally:
        formulas.close()
        cached.close()


def check_layout(external, baseline):
    if external.get('layout') != baseline.get('layout'):
        raise StructureChanged('检测到工作表名称、顺序或原表行列范围变化。请先到导入与导出窗口核对新结构，自动同步已暂停。')
    outside, previous = external.get('outside', {}), baseline.get('outside', {})
    if any(not same(outside.get(key), previous.get(key)) for key in set(outside) | set(previous)):
        raise StructureChanged('检测到原表范围外新增或删除的内容。请先通过导入窗口确认，自动同步已暂停。')


def compare(baseline, external, live):
    changes, conflicts = [], []
    before, excel, system = baseline['entries'], external['entries'], live['entries']
    for key in sorted(set(before) | set(excel) | set(system)):
        base, incoming, current = before.get(key), excel.get(key), system.get(key)
        external_changed, system_changed = not same(incoming, base), not same(current, base)
        if external_changed and system_changed and not same(incoming, current):
            sheet, row, col = json.loads(key)
            conflicts.append({'sheet': sheet, 'cell': wb.letter(col) + str(row),
                              'baseline': base or entry(), 'external': incoming or entry(),
                              'system': current or entry()})
        elif external_changed and not same(incoming, current):
            sheet, row, col = json.loads(key)
            incoming = incoming or entry()
            changes.append({'sheet': sheet, 'row': row, 'col': col,
                            'input': input_value(incoming)})
    return changes, conflicts


def state(c, doc):
    row = c.execute('SELECT * FROM external_workbook_sync WHERE document_id=?', (doc,)).fetchone()
    return dict(row) if row else None


def update_state(c, doc, **values):
    allowed = {'enabled', 'status', 'message', 'baseline_json', 'file_sha', 'revision',
               'generation', 'conflicts_json', 'updated_at'}
    if not set(values).issubset(allowed):
        raise ValueError('无效同步状态字段。')
    values['updated_at'] = db.now_str()
    c.execute('UPDATE external_workbook_sync SET ' + ','.join(key + '=?' for key in values) + ' WHERE document_id=?',
              [*values.values(), doc])


def note_status(doc, expected_generation, status, message, pause=False, conflicts=None):
    with r.transaction() as c:
        current = state(c, doc)
        if not current or current['generation'] != expected_generation:
            return
        values = {'status': status, 'message': message}
        if pause:
            values['enabled'] = 0
        if conflicts is not None:
            values['conflicts_json'] = dump(conflicts)
        update_state(c, doc, **values)


def live_state(c, doc):
    record = wb.getdoc(c, doc)
    rows, cells, formulas, _index = wb.load(c, doc)
    revision = c.execute('SELECT revision FROM workbook_state WHERE document_id=?', (doc,)).fetchone()[0]
    return record, rows, snapshot(rows, cells, formulas), revision


def export_bytes(record, rows):
    return wb.patch_xlsx(enhanced.doc_path(record), rows).getvalue()


def exported_baseline(data, live):
    file_view = read_external(data, live['limits'])
    return {**live, 'layout': file_view['layout'], 'outside': file_view['outside']}


def synchronize(doc, force_stable=False):
    """One polling cycle. force_stable is for deterministic integration tests."""
    with document_lock(doc):
        with r.transaction() as c:
            current = state(c, doc)
            if not current or not current['enabled']:
                return
            record, rows, live, revision = live_state(c, doc)
        generation = current['generation']
        baseline = json.loads(current['baseline_json'])
        if not baseline.get('entries'):
            note_status(doc, generation, 'error', '同步基线缺失，请明确选择重新导出系统数据。', pause=True)
            return
        try:
            path = safe_path(record)
            data, sha = read_bytes(path)
            observed_sha, rounds = _observations.get(doc, ('', 0))
            rounds = rounds + 1 if observed_sha == sha else 1
            _observations[doc] = (sha, rounds)
            if sha != current['file_sha'] and rounds < 2 and not force_stable:
                note_status(doc, generation, 'waiting_file', '已发现 Excel 保存，等待连续两次检测内容稳定。')
                return
            if sha == current['file_sha']:
                external = baseline
            else:
                external = read_external(data, baseline['limits'])
                check_layout(external, baseline)
            changes, conflicts = compare(baseline, external, live)
            if conflicts:
                note_status(doc, generation, 'conflict', '双方对同一单元格作了不同修改。请核对后明确选择恢复方式，本次未覆盖任何一方。', pause=True, conflicts=conflicts)
                return
            if changes:
                # Verify the observed file again before accepting its edits.
                if read_bytes(path)[1] != sha:
                    raise FileChanged('Excel 文件仍在保存，等待下一次稳定检测。')
                with r.transaction() as c:
                    latest = state(c, doc)
                    if not latest or not latest['enabled'] or latest['generation'] != generation:
                        return
                    result = wb.apply_changes(c, doc, changes, revision, origin='外部Excel保存')
                    record, rows, live, revision = live_state(c, doc)
                    # Accept the observed external state even if Excel later prevents
                    # the write-back. Pending system edits still differ from it.
                    update_state(c, doc, baseline_json=dump(external), file_sha=sha,
                                 revision=revision, status='waiting_write',
                                 message='Excel 修改已同步到系统，正在准备回写计算结果。', conflicts_json='[]')
                    baseline = external
            system_diff = any(not same(live['entries'].get(key), baseline['entries'].get(key))
                              for key in set(live['entries']) | set(baseline['entries']))
            limits_diff = live['limits'] != baseline.get('limits')
            if not changes and not system_diff and not limits_diff and current['status'] != 'waiting_write':
                with r.transaction() as c:
                    latest = state(c, doc)
                    if latest and latest['generation'] == generation and latest['enabled']:
                        update_state(c, doc, baseline_json=dump(external), file_sha=sha,
                                     revision=revision, status='active', message='原表与系统已同步。', conflicts_json='[]')
                return
            output = export_bytes(record, rows)
            next_baseline = exported_baseline(output, live)
            atomic_write(path, output, sha)
            output_sha = hashlib.sha256(output).hexdigest()
            with r.transaction() as c:
                latest = state(c, doc)
                if latest and latest['generation'] == generation and latest['enabled']:
                    update_state(c, doc, baseline_json=dump(next_baseline), file_sha=output_sha,
                                 revision=revision, status='active', message='原表与系统已同步。', conflicts_json='[]')
            _observations[doc] = (output_sha, 2)
        except StructureChanged as error:
            note_status(doc, generation, 'structure', str(error), pause=True)
        except FileChanged as error:
            note_status(doc, generation, 'waiting_file', str(error))
        except (PermissionError, FileNotFoundError, OSError) as error:
            note_status(doc, generation, 'waiting_write', '联动文件正在被 Excel 占用、尚未保存或暂时无法读写。保留待同步内容，稍后重试。' + str(error))
        except r.Conflict as error:
            note_status(doc, generation, 'waiting_file', str(error) + ' 下一轮重新核对。')
        except Exception as error:
            note_status(doc, generation, 'error', str(error) + ' 同步已暂停，未继续覆盖。', pause=True)


def authorize_operator():
    user = auth.current_user()
    if not user or user['role'] not in ('admin', 'bookkeeper'):
        abort(403)


@bp.get('/external-sync')
@auth.login_required
def page():
    docs = enhanced.documents()
    selected = request.args.get('doc')
    doc = next((item for item in docs if str(item['id']) == selected), None) if selected else (docs[0] if docs else None)
    if selected and doc is None:
        abort(404)
    current, conflicts, path, resolve_token = None, [], '', ''
    if doc:
        with r.transaction() as c:
            current = state(c, doc['id'])
        path = str(safe_path(doc))
        if current:
            conflicts = json.loads(current['conflicts_json'])
            if current['status'] in ('conflict', 'structure', 'error'):
                with r.transaction() as c:
                    _record, _rows, live, revision = live_state(c, doc['id'])
                sha = None
                try:
                    data, sha = read_bytes(pathlib.Path(path))
                    if current['status'] == 'conflict':
                        baseline = json.loads(current['baseline_json'])
                        external = read_external(data, baseline['limits'])
                        _changes, conflicts = compare(baseline, external, live)
                except Exception:
                    current['message'] = '联动文件暂时无法完整读取。可先保存并关闭 Excel；若文件已损坏，明确选择重新导出系统数据。'
                resolve_token = r.signer().dumps({'doc': doc['id'], 'generation': current['generation'],
                                                'revision': revision, 'file_sha': sha})
    user = auth.current_user()
    return render_template('external_sync.html', docs=docs, doc=doc, sync_state=current,
                           sync_path=path, conflicts=conflicts, status_labels=STATUS_LABELS,
                           can_control=user['role'] in ('admin', 'bookkeeper'), resolve_token=resolve_token)


@bp.post('/external-sync/<int:doc>/enable')
@auth.login_required
def enable(doc):
    authorize_operator()
    try:
        with document_lock(doc):
            with r.transaction() as c:
                record, rows, live, revision = live_state(c, doc)
                current = state(c, doc)
            path = safe_path(record)
            if current and current['baseline_json'] != '{}' and path.exists():
                if current['status'] in ('conflict', 'structure', 'error'):
                    raise ValueError('请先核对问题，并明确选择恢复方式。')
                with r.transaction() as c:
                    update_state(c, doc, enabled=1, status='active',
                                 generation=current['generation'] + 1,
                                 message='已开启自动检测，继续按上次同步基线核对。')
                    r.audit(c, '开启外部Excel同步', 'source_documents', doc, None, {'路径': str(path)})
            else:
                if path.exists():
                    with r.transaction() as c:
                        c.execute('INSERT OR IGNORE INTO external_workbook_sync(document_id) VALUES(?)', (doc,))
                        update_state(c, doc, enabled=0, status='error', message='联动目录已有同名文件，未自动覆盖。请核对并明确选择重新导出系统数据。')
                    raise ValueError('联动目录已有同名文件，未自动覆盖。请先核对并选择重新导出系统数据。')
                output = export_bytes(record, rows)
                baseline = exported_baseline(output, live)
                atomic_write(path, output, None)
                with r.transaction() as c:
                    c.execute('INSERT OR IGNORE INTO external_workbook_sync(document_id) VALUES(?)', (doc,))
                    latest = state(c, doc)
                    update_state(c, doc, enabled=1, status='active', message='联动工作副本已生成，自动同步已开启。',
                                 baseline_json=dump(baseline), file_sha=hashlib.sha256(output).hexdigest(),
                                 revision=revision, generation=latest['generation'] + 1, conflicts_json='[]')
                    r.audit(c, '开启外部Excel同步', 'source_documents', doc, None, {'路径': str(path)})
            _observations.pop(doc, None)
        flash('外部 Excel 自动同步已开启。请打开页面指定的联动文件。', 'ok')
    except Exception as error:
        flash(str(error), 'error')
    return redirect(url_for('external_sync.page', doc=doc))


@bp.post('/external-sync/<int:doc>/pause')
@auth.login_required
def pause(doc):
    authorize_operator()
    with document_lock(doc), r.transaction() as c:
        wb.getdoc(c, doc)
        current = state(c, doc)
        if current:
            # Keep a conflict/error state and its evidence while switching off.
            values = {'enabled': 0, 'generation': current['generation'] + 1}
            if current['status'] not in ('conflict', 'structure', 'error'):
                values.update(status='paused', message='已暂停自动同步，尚未同步的内容保持原样。')
            update_state(c, doc, **values)
            r.audit(c, '暂停外部Excel同步', 'source_documents', doc, None, {})
    flash('自动同步已暂停。', 'ok')
    return redirect(url_for('external_sync.page', doc=doc))


@bp.post('/external-sync/<int:doc>/resolve')
@auth.login_required
def resolve(doc):
    authorize_operator()
    mode = request.form.get('mode')
    if mode not in ('external', 'system') or request.form.get('confirmed') != '1':
        flash('请选择采用哪一方，并勾选确认。没有进行覆盖。', 'error')
        return redirect(url_for('external_sync.page', doc=doc))
    try:
        reviewed = r.signer().loads(request.form.get('resolve_token', ''))
        with document_lock(doc):
            with r.transaction() as c:
                record, rows, live, revision = live_state(c, doc)
                current = state(c, doc)
            path = safe_path(record)
            data, sha = read_bytes(path) if path.exists() else (None, None)
            if not current or reviewed != {'doc': doc, 'generation': current['generation'],
                                          'revision': revision, 'file_sha': sha}:
                raise r.Conflict('系统数据或 Excel 文件已变化，请重新打开同步页面核对后选择。')
            backup = backup_copy(path, data) if data is not None else ''
            if mode == 'external':
                if not current or not data:
                    raise ValueError('缺少原同步基线或联动文件，不能直接采用 Excel。')
                baseline = json.loads(current['baseline_json'])
                external = read_external(data, baseline['limits'])
                check_layout(external, baseline)
                changes = []
                for key, incoming in external['entries'].items():
                    if same(incoming, baseline['entries'].get(key)) or same(incoming, live['entries'].get(key)):
                        continue
                    sheet, row, col = json.loads(key)
                    changes.append({'sheet': sheet, 'row': row, 'col': col, 'input': input_value(incoming)})
                if read_bytes(path)[1] != sha:
                    raise FileChanged('Excel 文件再次保存，请重新查看后选择。')
                with r.transaction() as c:
                    if changes:
                        wb.apply_changes(c, doc, changes, revision, origin='操作员确认采用外部Excel')
                    record, rows, live, revision = live_state(c, doc)
                    latest = state(c, doc)
                    update_state(c, doc, enabled=1, status='waiting_write',
                                 baseline_json=dump(external), file_sha=sha, revision=revision,
                                 generation=latest['generation'] + 1, conflicts_json='[]',
                                 message='已按操作员选择采用 Excel 修改，等待回写合并结果。')
                    r.audit(c, '确认外部Excel冲突处理', 'source_documents', doc, None, {'采用': 'Excel修改', '备份': backup})
            else:
                output = export_bytes(record, rows)
                baseline = exported_baseline(output, live)
                atomic_write(path, output, sha)
                with r.transaction() as c:
                    c.execute('INSERT OR IGNORE INTO external_workbook_sync(document_id) VALUES(?)', (doc,))
                    latest = state(c, doc)
                    update_state(c, doc, enabled=1, status='active', baseline_json=dump(baseline),
                                 file_sha=hashlib.sha256(output).hexdigest(), revision=revision,
                                 generation=latest['generation'] + 1, conflicts_json='[]',
                                 message='已按操作员选择重新导出系统数据。原联动文件已备份。')
                    r.audit(c, '确认外部Excel冲突处理', 'source_documents', doc, None, {'采用': '系统数据', '备份': backup})
            _observations.pop(doc, None)
        flash('已执行所选恢复方式。处理前的联动文件保存在同一目录。', 'ok')
    except Exception as error:
        flash(str(error) + ' 未继续自动覆盖。', 'error')
    return redirect(url_for('external_sync.page', doc=doc))


def worker(app):
    while True:
        try:
            with app.app_context(), app.test_request_context('/__external_excel_worker'):
                documents = db.query('SELECT document_id FROM external_workbook_sync WHERE enabled=1')
                for item in documents:
                    try:
                        synchronize(item['document_id'])
                    except Exception:
                        app.logger.exception('外部Excel同步检测失败')
        except Exception:
            app.logger.exception('外部Excel同步任务失败')
        threading.Event().wait(2)


def configure(app):
    with r.transaction() as c:
        c.executescript(SCHEMA)
    app.register_blueprint(bp)
    if os.environ.get('LEDGER_DISABLE_EXTERNAL_THREAD', '').lower() in ('1', 'true', 'yes'):
        return
    if id(app) not in _started:
        _started.add(id(app))
        threading.Thread(target=worker, args=(app,), daemon=True, name='external-excel-sync').start()
