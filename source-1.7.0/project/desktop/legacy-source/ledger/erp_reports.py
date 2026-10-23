"""As-of ERP reports and safe XLSX exports, without legacy recalculation.

Balances come from bill_rows' immutable dated postings. Source date filters
select bills, rather than cutting off their earlier allocations. Money stays
integer cents; workbook amounts are exact two-decimal text for Excel precision.
"""
import datetime as dt
import re
from copy import deepcopy
from io import BytesIO
from erp_values import ErpError, date, identifier, permitted, require, today

PARAMETERS = frozenset(('asOf', 'from', 'to', 'partyId', 'kind', 'status', 'creatorId'))
KINDS = frozenset(('purchase', 'receipt', 'receivable', 'payable', 'opening_receivable',
                   'opening_payable', 'settlement', 'reversal'))
STATUSES = frozenset(('draft', 'pending', 'confirmed', 'rejected', 'reversed'))
MAX_ROWS = 50000
BUCKETS = (('not_due', '未到期（含当日到期）'), ('1_30', '逾期 1–30 天'),
           ('31_60', '逾期 31–60 天'), ('61_90', '逾期 61–90 天'), ('91_plus', '逾期超过 90 天'))


def parse_filters(params):
    if not hasattr(params, 'items') or set(params) - PARAMETERS:
        raise ErpError('报表筛选字段无效。')
    result = {}
    for key, value in params.items():
        if hasattr(params, 'getlist') and len(params.getlist(key)) != 1:
            raise ErpError('报表筛选字段不能重复。')
        if not isinstance(value, str) or len(value) > 300 or re.search(r'[\x00-\x1f\x7f]', value):
            raise ErpError('报表筛选内容无效。')
        if value:
            result[key] = value
    result['asOf'] = date(result.get('asOf', today()))
    for key in ('from', 'to'):
        if key in result:
            result[key] = date(result[key])
    if result.get('from', '0001-01-01') > result.get('to', result['asOf']) or result.get('from', '0001-01-01') > result['asOf']:
        raise ErpError('报表起始日期晚于结束日期或截至日期。')
    if 'kind' in result and result['kind'] not in KINDS:
        raise ErpError('报表单据类型无效。')
    if 'status' in result and result['status'] not in STATUSES:
        raise ErpError('报表单据状态无效。')
    party = result.get('partyId')
    if party:
        if party.startswith('CUS:'):
            identifier(party[4:])
        elif not re.fullmatch(r'SUP:[1-9][0-9]{0,18}', party):
            raise ErpError('报表往来单位标识无效。')
    if 'creatorId' in result and not re.fullmatch(r'[A-Za-z0-9_-]{1,100}', result['creatorId']):
        raise ErpError('报表制单人标识无效。')
    return result


def _cents(value):
    if type(value) is not int:
        raise ErpError('报表金额不是整数分，停止生成。', 500)
    return value


def _eligible(row, filters, source=None):
    source = source or row
    business_date = date(row['date'])
    upper = min(filters['asOf'], filters.get('to', filters['asOf']))
    if business_date > upper or business_date < filters.get('from', '0001-01-01'):
        return False
    if filters.get('partyId') and row['partyId'] != filters['partyId']:
        return False
    if filters.get('creatorId') and source.get('authorId') != filters['creatorId']:
        return False
    if filters.get('status') and source.get('status') != filters['status']:
        return False
    if filters.get('kind'):
        key = filters['kind']
        if key in ('payable',) and row.get('kind') == key:
            return True
        if source.get('kind') != key:
            return False
    return True


def _sensitive(document, documents):
    if document.get('redacted'):
        return True
    kind, payload = document['kind'], document.get('payload', {})
    if kind in ('purchase', 'receipt', 'opening_payable') or kind == 'settlement' and payload.get('direction') == 'pay':
        return True
    if kind == 'reversal':
        target = documents.get(payload.get('targetId'))
        # A missing target must never become an accidental sensitive bypass.
        return target is None or target['kind'] == 'reversal' or _sensitive(target, documents)
    return False


def build_report(actor, params, bills, documents):
    """Pure report calculation; caller supplies bill postings already as-of."""
    require(actor, 'erp.reports.view')
    filters = parse_filters(params)
    if len(bills) > MAX_ROWS or len(documents) > MAX_ROWS:
        raise ErpError('报表超过五万条，请缩小筛选范围。', 413)
    sensitive = permitted(actor, 'erp.reports.sensitive')
    documents = {row['id']: deepcopy(row) for row in documents}
    reversals = {row.get('payload', {}).get('targetId'): row
                 for row in documents.values()
                 if row['kind'] == 'reversal' and row.get('status') in ('confirmed', 'reversed') and row['date'] <= filters['asOf']}
    for row in documents.values():
        if row['kind'] != 'reversal' and row.get('status') in ('confirmed', 'reversed'):
            row['status'] = 'reversed' if row['id'] in reversals else 'confirmed'
    rows, cash, purchase = [], [], []
    for bill in bills:
        if bill['kind'] == 'payable' and not sensitive:
            continue
        source = documents.get(bill['id'])
        if source is None:
            raise ErpError('往来余额缺少来源单据，停止生成。', 500)
        if not _eligible(bill, filters, source):
            continue
        row = deepcopy(bill)
        for key in ('amountCents', 'settledCents', 'balanceCents'):
            _cents(row[key])
        if row['balanceCents'] != row['amountCents'] - row['settledCents'] or min(row['amountCents'], row['settledCents'], row['balanceCents']) < 0:
            raise ErpError('往来分录余额不一致，停止生成。', 500)
        row.update(status=source['status'], sourceKind=source['kind'], authorId=source.get('authorId'), authorName=source.get('authorName'))
        rows.append(row)
    for document in documents.values():
        if not sensitive and _sensitive(document, documents):
            continue
        if not _eligible(document, filters):
            continue
        if document['kind'] in ('purchase', 'receipt') and sensitive:
            purchase.append(document)
        if document.get('status') not in ('confirmed', 'reversed'):
            continue
        source, sign = document, 1
        if document['kind'] == 'reversal':
            source = documents.get(document.get('payload', {}).get('targetId'))
            if source is None or source['kind'] != 'settlement':
                continue
            sign = -1
        elif document['kind'] != 'settlement':
            continue
        payload = source.get('payload', {})
        if payload.get('direction') not in ('receive', 'pay'):
            raise ErpError('收付款方向不完整，停止生成。', 500)
        if payload['direction'] == 'pay' and not sensitive:
            continue
        cash.append({'id': document['id'], 'number': document['number'], 'kind': document['kind'],
                     'partyId': document['partyId'], 'partyName': document['partyName'], 'date': document['date'],
                     'direction': payload['direction'], 'account': payload.get('account', ''),
                     'amountCents': sign * _cents(source['amountCents']), 'status': document['status'],
                     'note': document.get('note', ''), 'authorId': document.get('authorId'),
                     'authorName': document.get('authorName'), 'targetId': source['id'] if sign < 0 else None})
    for collection in (rows, cash, purchase):
        collection.sort(key=lambda row: (row['date'], row['id']))
    totals = {'receivableCents': sum(row['balanceCents'] for row in rows if row['kind'] == 'receivable'),
              'receivableCount': sum(1 for row in rows if row['kind'] == 'receivable' and row['balanceCents'] > 0),
              'receivedCents': sum(row['amountCents'] for row in cash if row['direction'] == 'receive'),
              'receivedCount': sum(1 for row in cash if row['direction'] == 'receive'), 'rowCount': len(rows)}
    if sensitive:
        totals.update(payableCents=sum(row['balanceCents'] for row in rows if row['kind'] == 'payable'),
                      payableCount=sum(1 for row in rows if row['kind'] == 'payable' and row['balanceCents'] > 0),
                      paidCents=sum(row['amountCents'] for row in cash if row['direction'] == 'pay'),
                      paidCount=sum(1 for row in cash if row['direction'] == 'pay'), purchaseCount=len(purchase))
    aging = [{'key': key, 'label': label, 'receivableCents': 0, 'receivableCount': 0,
              **({'payableCents': 0, 'payableCount': 0} if sensitive else {})} for key, label in BUCKETS]
    cutoff = dt.date.fromisoformat(filters['asOf'])
    for row in rows:
        if row['balanceCents'] == 0:
            continue
        days = (cutoff - dt.date.fromisoformat(date(row['dueDate']))).days
        index = 0 if days <= 0 else 1 if days <= 30 else 2 if days <= 60 else 3 if days <= 90 else 4
        aging[index][row['kind'] + 'Cents'] += row['balanceCents']
        aging[index][row['kind'] + 'Count'] += 1
    if any(abs(value)>9007199254740991 for key,value in totals.items() if key.endswith('Cents')):
        raise ErpError('汇总超过浏览器可安全显示的金额精度，请按往来单位或日期缩小范围。',413)
    return {'asOf': filters['asOf'], 'filters': filters, 'rows': rows, 'totals': totals,
            'aging': aging, 'cash': cash, 'purchase': purchase, 'sensitiveAccess': sensitive}


def report(actor, params):
    require(actor, 'erp.reports.view')
    filters = parse_filters(params)
    from erp_common import bill_rows, public_document
    import repairs
    with repairs.transaction() as connection:
        raw = connection.execute('SELECT * FROM erp_documents WHERE date<=? ORDER BY date,id LIMIT ?',
                                 (filters['asOf'], MAX_ROWS + 1)).fetchall()
        if len(raw) > MAX_ROWS:
            raise ErpError('报表来源超过五万条，暂不能安全导出完整结果。', 413)
        documents = [public_document(connection, dict(row), actor, as_of=filters['asOf']) for row in raw]
        bills = bill_rows(connection, actor, filters['asOf'])
        return build_report(actor, filters, bills, documents)


def _money_text(cents):
    cents = _cents(cents)
    negative = '-' if cents < 0 else ''
    whole, fraction = divmod(abs(cents), 100)
    return f'{negative}{whole}.{fraction:02d}'


def workbook(data):
    """Serialize a previously authorized report; every string cell stays text."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    book = Workbook()
    book.remove(book.active)
    def sheet(name, headings, values):
        output = book.create_sheet(name)
        for row_number, row in enumerate([headings, *values], 1):
            for column_number, value in enumerate(row, 1):
                cell = output.cell(row_number, column_number)
                if isinstance(value, str):
                    cell.value = value
                    cell.data_type = 's'
                else:
                    cell.value = value
                if row_number == 1:
                    cell.fill = PatternFill('solid', fgColor='23251F')
                    cell.font = Font(name='微软雅黑', color='FFFFFF', bold=True, size=11)
                else:
                    cell.font = Font(name='微软雅黑', size=11)
                cell.alignment = Alignment(vertical='top', wrap_text=True)
        output.freeze_panes = 'A2'
        output.auto_filter.ref = output.dimensions
        for cells in output.columns:
            width = min(42, max(14, max(len(str(cell.value or '')) for cell in cells) + 3))
            output.column_dimensions[cells[0].column_letter].width = width
        output.sheet_properties.pageSetUpPr.fitToPage = True
        output.page_setup.orientation = 'landscape'
        output.page_setup.paperSize = output.PAPERSIZE_A4
        output.page_setup.fitToWidth = 1
        output.page_setup.fitToHeight = 0
        output.print_title_rows = '1:1'
        return output
    sheet('往来余额', ['单号', '方向', '单位', '日期', '到期日', '原金额 / 元（精确文本）', '已核销 / 元（精确文本）', '未结 / 元（精确文本）', '来源类型', '状态', '制单人'],
          [[row['number'], '应付' if row['kind'] == 'payable' else '应收', row['partyName'], row['date'], row['dueDate'],
            *[_money_text(row[key]) for key in ('amountCents', 'settledCents', 'balanceCents')], row.get('sourceKind'), row.get('status'), row.get('authorName')]
           for row in data['rows']])
    labels = {'receivableCents': '未结应收 / 元', 'payableCents': '未结应付 / 元', 'receivedCents': '净确认收款 / 元', 'paidCents': '净确认付款 / 元',
              'receivableCount': '未结应收笔数', 'payableCount': '未结应付笔数', 'receivedCount': '收款及冲销记录数', 'paidCount': '付款及冲销记录数', 'rowCount': '往来明细记录数', 'purchaseCount': '采购验收记录数'}
    sheet('汇总', ['指标', '数值'], [[labels.get(key, key), _money_text(value) if key.endswith('Cents') else value] for key, value in data['totals'].items()])
    sensitive = data.get('sensitiveAccess', False)
    sheet('账龄', ['账龄', '应收 / 元（精确文本）', '应收笔数'] + (['应付 / 元（精确文本）', '应付笔数'] if sensitive else []),
          [[row['label'], _money_text(row['receivableCents']), row['receivableCount']] + ([_money_text(row['payableCents']), row['payableCount']] if sensitive else []) for row in data['aging']])
    sheet('收付款', ['单号', '日期', '方向', '单位', '账户', '金额 / 元（精确文本）', '备注', '制单人', '冲销原单ID'],
          [[row['number'], row['date'], '付款' if row['direction'] == 'pay' else '收款', row['partyName'], row['account'], _money_text(row['amountCents']), row.get('note'), row.get('authorName'), row.get('targetId')] for row in data['cash']])
    sheet('采购验收', ['单号', '类型', '单位', '日期', '金额 / 元（精确文本）', '状态', '制单人', '备注'],
          [[row['number'], row['kind'], row['partyName'], row['date'], _money_text(row['amountCents']), row['status'], row.get('authorName'), row.get('note')] for row in data['purchase']] if sensitive else [])
    sheet('筛选条件', ['条件', '值'], [[key, value] for key, value in data['filters'].items()] + [['金额口径', '整数分汇总；Excel金额为精确两位小数文本。'], ['截至日口径', '按业务日期读取确认、核销、冲销分录；日期筛选选择来源单据，不裁掉其历史核销。']])
    output = BytesIO()
    book.save(output)
    book.close()
    output.seek(0)
    return output


def export(actor, params):
    return workbook(report(actor, params))
