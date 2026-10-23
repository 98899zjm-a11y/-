"""Immutable bundled reference and explicit operator-approved ledger import."""
import json,pathlib,zipfile,xml.etree.ElementTree as ET
import openpyxl
from flask import Blueprint,abort,flash,redirect,render_template,request,send_file,url_for
from app import auth,db
import repairs as r
import enhancements as enhanced
import source_sync as sync

bp=Blueprint('bundled_source',__name__)

def asset_path():return pathlib.Path(__file__).resolve().parent/'assets'/'原版Excel.xlsx'

def register_if_empty(c):
    """Store reference metadata only. Business entries require a POST confirmation."""
    if c.execute('SELECT 1 FROM source_documents LIMIT 1').fetchone():return None
    path=asset_path()
    if not path.is_file():return None
    doc=enhanced.register_file(path,'原版Excel.xlsx',c)
    r.audit(c,'登记内置原版Excel','source_documents',doc,None,{'用途':'原表预览和确认导入','自动入账':False})
    return doc

@bp.get('/original-excel/download')
@auth.login_required
def download():
    path=asset_path()
    if not path.is_file():abort(404)
    return send_file(path,as_attachment=True,download_name='鑫金源珠宝_原版Excel.xlsx',mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

@bp.get('/original-excel/preview')
@auth.login_required
def preview():
    path=asset_path()
    if not path.is_file():abort(404)
    values=openpyxl.load_workbook(path,read_only=True,data_only=True)
    formulas=openpyxl.load_workbook(path,read_only=True,data_only=False)
    try:
        selected=request.args.get('sheet',values.sheetnames[0])
        if selected not in values.sheetnames:abort(404)
        ws=values[selected];fs=formulas[selected]
        try:start=int(request.args.get('start',1))
        except ValueError:start=1
        start=max(1,min(start,ws.max_row));end=min(ws.max_row,start+49)
        namespace={'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
        with zipfile.ZipFile(path) as book:
            xml=ET.fromstring(book.read('xl/worksheets/sheet'+str(values.sheetnames.index(selected)+1)+'.xml'))
        present=[c.get('r') for c in xml.findall('m:sheetData/m:row/m:c',namespace) if c.find('m:v',namespace) is not None or c.find('m:f',namespace) is not None or c.find('m:is',namespace) is not None]
        count=min(40,max([openpyxl.utils.cell.column_index_from_string(openpyxl.utils.cell.coordinate_from_string(x)[0]) for x in present] or [1]))
        rows=[]
        value_rows=ws.iter_rows(min_row=start,max_row=end,max_col=count)
        formula_rows=fs.iter_rows(min_row=start,max_row=end,max_col=count)
        for rn,(row,formula_row) in enumerate(zip(value_rows,formula_rows),start):
            shown=[]
            for cell,other in zip(row,formula_row):
                shown.append({'value':cell.value,'numeric':isinstance(cell.value,(int,float)) and not isinstance(cell.value,bool),'formula':other.value if other.data_type=='f' else ''})
            rows.append({'row':rn,'cells':shown})
        return render_template('bundled_excel.html',sheets=values.sheetnames,sheet=selected,rows=rows,columns=[openpyxl.utils.get_column_letter(i) for i in range(1,count+1)],start=start,end=end,max_row=ws.max_row)
    finally:values.close();formulas.close()

@bp.post('/original-excel/import')
@auth.login_required
def import_book():
    if auth.current_user()['role']!='admin':abort(403)
    try:
        if request.form.get('confirm')!='1':raise ValueError('请先确认将原表明细导入当前账本。')
        with r.transaction() as c:
            document=int(request.form.get('document_id','0'))
            record=c.execute('SELECT * FROM source_documents WHERE id=?',(document,)).fetchone()
            if not record:raise ValueError('原表不存在，请重新载入页面。')
            record=dict(record);r.check_token('source_documents',record,request.form.get('snapshot'))
            sheets=[x[0] for x in c.execute("SELECT DISTINCT sheet FROM source_rows WHERE document_id=? AND kind IN ('sales','stock','material')",(document,))]
            if not sheets:raise ValueError('原表没有识别到可导入类目，请到原表对账核对。')
            result=sync.import_document(document,c,selected=sheets)
            r.audit(c,'确认导入原表','source_documents',document,None,result)
        flash('原表已确认导入：'+json.dumps(result,ensure_ascii=False)+'。待核对项目可在原表对账中继续处理。','ok')
    except Exception as error:flash('本次未导入：'+str(error),'error')
    return redirect(url_for('sales.index'))

def configure(app):
    with app.test_request_context():
        with r.transaction() as c:register_if_empty(c)
    app.register_blueprint(bp)
    @app.context_processor
    def original_reference():
        doc=db.query('SELECT * FROM source_documents ORDER BY id DESC LIMIT 1',one=True)
        return {'original_excel_available':asset_path().is_file(),'original_import_doc':dict(doc) if doc else None,'original_import_token':r.row_token('source_documents',doc) if doc else ''}
